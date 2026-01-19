#!/usr/bin/env python3
"""
Script mejorado para convertir Markdown a Word (.docx)
Requisitos: 
- Times New Roman 12 para texto normal
- Estilos de Título para Tabla de Contenido
- Soporte para tildes y eñes
"""
import re
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_font(run, name='Times New Roman', size=12, color=None, bold=False):
    run.font.name = name
    # Solución para problema de visualización de fuentes en Word (obligar a usar la fuente)
    r = run._element
    r.rPr.rFonts.set(qn('w:eastAsia'), name)
    r.rPr.rFonts.set(qn('w:ascii'), name)
    r.rPr.rFonts.set(qn('w:hAnsi'), name)
    
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)

def convert_markdown_to_docx(md_file, docx_file):
    if not os.path.exists(md_file):
        print(f"Error: No se encuentra el archivo {md_file}")
        return

    # Leer el archivo markdown con codificación UTF-8
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # Crear documento Word
    doc = Document()
    
    # Configurar estilo Normal (Times New Roman 12)
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    # Procesar línea por línea
    lines = content.split('\n')
    in_code_block = False
    code_lines = []
    
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        
        # Bloques de código
        if line.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lines = []
            else:
                in_code_block = False
                if code_lines:
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Inches(0.5)
                    p.paragraph_format.space_before = Pt(6)
                    p.paragraph_format.space_after = Pt(6)
                    
                    # Sombreado de fondo
                    shading_elm = OxmlElement('w:shd')
                    shading_elm.set(qn('w:fill'), 'F0F0F0')
                    p._p.get_or_add_pPr().append(shading_elm)
                    
                    run = p.add_run('\n'.join(code_lines))
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9)
                code_lines = []
            i += 1
            continue
        
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue
        
        # Líneas vacías
        if not line.strip():
            doc.add_paragraph()
            i += 1
            continue
        
        # Encabezados (Títulos para TOC)
        if line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            text = line.lstrip('#').strip()
            
            # Usar estilos nativos de Word para Tabla de Contenido
            style_name = f'Heading {level}'
            if style_name not in doc.styles:
                style_name = 'Heading 1' # Fallback
            
            p = doc.add_paragraph(text, style=style_name)
            
            # Personalizar apariencia pero manteniendo el nivel
            run = p.runs[0]
            if level == 1:
                set_font(run, name='Arial', size=18, color=(0, 51, 102), bold=True)
            elif level == 2:
                set_font(run, name='Arial', size=14, color=(0, 102, 204), bold=True)
            else:
                set_font(run, name='Arial', size=12, color=(51, 102, 153), bold=True)
            
            i += 1
            continue
        
        # Listas
        if line.strip().startswith(('-', '*')):
            text = line.strip()[1:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.left_indent = Inches(0.25)
            
            # Procesar formato dentro de la lista (negritas, etc)
            parts = re.split(r'(\*\*.*?\*\*|`.*?`)', text)
            for part in parts:
                if not part: continue
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    set_font(run, bold=True)
                elif part.startswith('`') and part.endswith('`'):
                    run = p.add_run(part[1:-1])
                    run.font.name = 'Consolas'
                    run.font.size = Pt(10)
                else:
                    run = p.add_run(part)
                    set_font(run)
            i += 1
            continue

        # Texto normal con formato (Times New Roman 12)
        p = doc.add_paragraph()
        parts = re.split(r'(\*\*.*?\*\*|`.*?`)', line)
        
        for part in parts:
            if not part: continue
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                set_font(run, bold=True)
            elif part.startswith('`') and part.endswith('`'):
                run = p.add_run(part[1:-1])
                run.font.name = 'Consolas'
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(199, 37, 78)
            else:
                run = p.add_run(part)
                set_font(run) # Aplica Times New Roman 12
        
        i += 1
    
    # Guardar documento
    doc.save(docx_file)
    print(f"Documento generado exitosamente: {docx_file}")

if __name__ == '__main__':
    # Usar el nuevo archivo completo
    convert_markdown_to_docx('API_DOCUMENTATION_COMPREHENSIVE.md', 'API_CORE_IDENTITY_DETALLADA.docx')
    # También generar el anterior si se desea
    # convert_markdown_to_docx('API_ENDPOINTS.md', 'API_ENDPOINTS.docx')
